import csv
import io
import json
import os
from decimal import Decimal
from pathlib import Path

import openpyxl
import psycopg
import pytest
from fastapi import HTTPException

from api import _paper
from api.ledger import export_ledger, ledger, router as ledger_router
from api._paper import handler


ROOT = Path(__file__).resolve().parents[1]


def _run_sql_file(conn, path):
    conn.execute(Path(path).read_text(encoding="utf-8"))


@pytest.fixture()
def clean_paper_db(monkeypatch):
    dsn = os.environ.get("DATABASE_URL")
    if not dsn:
        pytest.skip("DATABASE_URL is required for PostgreSQL paper-ledger integration tests")
    with psycopg.connect(dsn) as conn:
        conn.execute("DROP SCHEMA public CASCADE")
        conn.execute("CREATE SCHEMA public")
        _run_sql_file(conn, ROOT / "db/migrations/001_paper_base.sql")
        _run_sql_file(conn, ROOT / "db/migrations/002_paper_trade_ledger.sql")
        conn.commit()
        fk_tables = {
            "fk_paper_orders_account": "zerqen_paper_orders",
            "fk_paper_fills_account": "zerqen_paper_fills",
            "fk_paper_fills_order": "zerqen_paper_fills",
            "fk_paper_positions_account": "zerqen_paper_positions",
            "fk_paper_positions_order": "zerqen_paper_positions",
            "fk_paper_positions_fill": "zerqen_paper_positions",
            "fk_paper_snapshots_account": "zerqen_paper_equity_snapshots",
            "fk_paper_events_account": "zerqen_paper_events",
            "fk_paper_decisions_account": "zerqen_paper_decisions",
            "fk_paper_decisions_order": "zerqen_paper_decisions",
            "fk_paper_trades_account": "zerqen_paper_trades",
            "fk_paper_trades_decision": "zerqen_paper_trades",
            "fk_paper_trades_entry_order": "zerqen_paper_trades",
            "fk_paper_trades_exit_order": "zerqen_paper_trades",
            "fk_paper_trades_entry_fill": "zerqen_paper_trades",
            "fk_paper_trades_exit_fill": "zerqen_paper_trades",
            "fk_paper_trades_snapshot": "zerqen_paper_trades",
            "fk_paper_trades_audit": "zerqen_paper_trades",
        }
        for name, table in fk_tables.items():
            conn.execute(f'ALTER TABLE "{table}" VALIDATE CONSTRAINT "{name}"')
        conn.commit()
    monkeypatch.setenv("ZERQEN_DASHBOARD_TOKEN", "test-token")
    monkeypatch.setenv("ZERQEN_VAULT_KEY", "test-vault-key")
    with _paper.db() as conn:
        assert conn.execute("SELECT 1").fetchone()[0] == 1
    return dsn


def _request(monkeypatch, action, **extra):
    captured = []

    def capture_send(_handler, status, payload):
        captured.append((status, payload))

    monkeypatch.setattr(_paper, "send", capture_send)
    body = json.dumps({"action": action, **extra}).encode()
    h = object.__new__(handler)
    h.headers = {
        "x-zerqen-dashboard-token": "test-token",
        "Content-Length": str(len(body)),
    }
    h.rfile = io.BytesIO(body)
    handler.do_POST(h)
    assert captured
    return captured[-1]


def _deterministic_market(monkeypatch):
    state = {"price": Decimal(100)}

    def fake_market(exchange_id, symbol, timeframe="1h", limit=120):
        del exchange_id, symbol, timeframe
        price = state["price"]
        rows = []
        for i in range(max(limit, 60)):
            ts = 1_790_000_000_000 + i * 3_600_000
            rows.append([ts, price, price + Decimal(1), price - Decimal(1), price, Decimal(1)])
        return price, rows

    monkeypatch.setattr(_paper, "fetch_public_market", fake_market)
    return state


def test_ledger_route_is_registered_and_authentication_is_enforced(clean_paper_db):
    assert any(getattr(route, "path", None) == "/api/ledger" for route in ledger_router.routes)
    index_source = (ROOT / "api/index.py").read_text(encoding="utf-8")
    assert "app.include_router(ledger_router)" in index_source
    with pytest.raises(HTTPException) as exc:
        ledger(x_zerqen_dashboard_token="wrong-token")
    assert exc.value.status_code == 401


def test_clean_migrations_are_idempotent_and_have_trace_constraints(clean_paper_db):
    with psycopg.connect(clean_paper_db) as conn:
        _run_sql_file(conn, ROOT / "db/migrations/001_paper_base.sql")
        _run_sql_file(conn, ROOT / "db/migrations/002_paper_trade_ledger.sql")
        conn.commit()
        tables = {
            row[0]
            for row in conn.execute(
                "SELECT tablename FROM pg_tables WHERE schemaname='public'"
            ).fetchall()
        }
        assert "zerqen_paper_trades" in tables
        assert "zerqen_paper_decisions" in tables
        assert "zerqen_paper_equity_snapshots" in tables
        assert "zerqen_paper_events" in tables
        fks = {
            row[0]
            for row in conn.execute(
                "SELECT conname FROM pg_constraint WHERE contype='f' AND conname LIKE 'fk_paper_%'"
            ).fetchall()
        }
        assert {
            "fk_paper_trades_decision",
            "fk_paper_trades_entry_order",
            "fk_paper_trades_exit_order",
            "fk_paper_trades_entry_fill",
            "fk_paper_trades_exit_fill",
            "fk_paper_trades_snapshot",
            "fk_paper_trades_audit",
        } <= fks
        assert conn.execute(
            "SELECT 1 FROM pg_indexes WHERE indexname='idx_paper_decisions_account_signal'"
        ).fetchone()
        assert conn.execute(
            "SELECT 1 FROM pg_indexes WHERE indexname='idx_paper_positions_position_id'"
        ).fetchone()


def test_deterministic_paper_lifecycle_reconciles_ledger_compounding_and_exports(
    clean_paper_db, monkeypatch
):
    market = _deterministic_market(monkeypatch)

    status, payload = _request(
        monkeypatch,
        "initialize",
        starting_capital="10000",
        exchange_id="okx",
        symbol="BTC/USDT",
        timeframe="1h",
    )
    assert status == 200
    assert payload["live_trading"] is False

    status, payload = _request(
        monkeypatch, "test_order", symbol="BTC/USDT", side="buy", quantity="10"
    )
    assert status == 200
    assert payload["live_trading"] is False

    with psycopg.connect(clean_paper_db) as conn:
        position = conn.execute(
            "SELECT position_id,quantity,average_entry,risk_at_entry,opening_equity,entry_order_id,entry_fill_id "
            "FROM zerqen_paper_positions WHERE account_id='default' AND symbol='BTC/USDT'"
        ).fetchone()
        assert position
        entry_order_id, entry_fill_id = position[5], position[6]

    market["price"] = Decimal(110)
    status, payload = _request(
        monkeypatch, "test_order", symbol="BTC/USDT", side="sell", quantity="10"
    )
    assert status == 200
    assert not payload["positions"]

    with psycopg.connect(clean_paper_db) as conn:
        trade = conn.execute(
            """SELECT trade_id,decision_id,position_id,equity_snapshot_id,audit_event_id,
                      signal_id,entry_order_id,exit_order_id,entry_fill_id,exit_fill_id,
                      entry_price,exit_price,quantity,risk_at_entry,gross_pnl,fees,slippage,
                      funding,net_pnl,r_multiple,opening_equity,closing_equity
               FROM zerqen_paper_trades WHERE account_id='default'"""
        ).fetchone()
        state = conn.execute(
            "SELECT starting_equity,realized_pnl,fees,slippage,cash FROM zerqen_paper_state WHERE account_id='default'"
        ).fetchone()
        snapshot = conn.execute(
            """SELECT id,equity,realized_pnl,unrealized_pnl,drawdown,daily_pnl
               FROM zerqen_paper_equity_snapshots
               WHERE account_id='default' ORDER BY id DESC LIMIT 1"""
        ).fetchone()
        fills = conn.execute(
            """SELECT fill_id,client_order_id,requested_price,price,quantity,fee,slippage
               FROM zerqen_paper_fills WHERE account_id='default' ORDER BY created_at"""
        ).fetchall()
        decisions = conn.execute(
            "SELECT decision_id,order_id,rejected FROM zerqen_paper_decisions WHERE account_id='default' ORDER BY signal_timestamp"
        ).fetchall()
        audit = conn.execute(
            "SELECT event_id,event_type,payload FROM zerqen_paper_events WHERE account_id='default' ORDER BY created_at"
        ).fetchall()

    assert trade is not None
    assert len(fills) == 2
    assert len(decisions) == 2
    assert len(audit) >= 3

    (
        _trade_id, decision_id, position_id, snapshot_id, audit_id, signal_id,
        trade_entry_order, exit_order_id, trade_entry_fill, exit_fill_id,
        entry_actual, exit_actual, qty, _risk, gross, fees, slippage, funding,
        net, r_multiple, opening_equity, closing_equity,
    ) = trade

    exit_fill = fills[1]
    assert trade_entry_order == entry_order_id
    assert trade_entry_fill == entry_fill_id
    assert exit_order_id == exit_fill[1]
    assert exit_fill_id == exit_fill[0]
    assert decision_id == decisions[-1][0]
    assert position_id == position[0]
    assert snapshot_id == snapshot[0]
    assert audit_id in {row[0] for row in audit}
    assert signal_id.startswith("signal-")

    requested_entry = Decimal(100)
    requested_exit = Decimal(110)
    quantity = Decimal(10)
    expected_entry_actual = requested_entry + requested_entry * Decimal("0.0005")
    expected_exit_actual = requested_exit - requested_exit * Decimal("0.0005")
    expected_entry_fee = expected_entry_actual * quantity * Decimal("0.001")
    expected_exit_fee = expected_exit_actual * quantity * Decimal("0.001")
    expected_slippage = (expected_entry_actual-requested_entry)*quantity + (requested_exit-expected_exit_actual)*quantity
    expected_gross = (requested_exit-requested_entry)*quantity
    expected_net = expected_gross - expected_entry_fee - expected_exit_fee - expected_slippage
    expected_closing_equity = Decimal(10000) + expected_net

    assert Decimal(str(entry_actual)) == expected_entry_actual
    assert Decimal(str(exit_actual)) == expected_exit_actual
    assert Decimal(str(qty)) == quantity
    assert Decimal(str(gross)) == expected_gross
    assert Decimal(str(fees)) == expected_entry_fee + expected_exit_fee
    assert Decimal(str(slippage)) == expected_slippage
    assert Decimal(str(funding)) == Decimal(0)
    assert Decimal(str(net)) == expected_net
    assert Decimal(str(state[1])) == expected_net
    assert Decimal(str(snapshot[1])) == expected_closing_equity
    assert Decimal(str(snapshot[2])) == expected_net
    assert Decimal(str(snapshot[3])) == Decimal(0)
    assert Decimal(str(snapshot[5])) == expected_net
    assert Decimal(str(opening_equity)) == Decimal(10000)
    assert Decimal(str(closing_equity)) == expected_closing_equity
    assert Decimal(str(r_multiple)).quantize(Decimal("0.000001")) == (
        expected_net / Decimal(50)
    ).quantize(Decimal("0.000001"))

    api_data = ledger(x_zerqen_dashboard_token="test-token")
    assert api_data["account"]["realized_pnl"] == float(expected_net)
    assert api_data["account"]["net_pnl"] == float(expected_net)
    assert api_data["account"]["actual_return"] == float(expected_net / Decimal(10000))
    assert len(api_data["trades"]) == 1
    api_trade = api_data["trades"][0]
    assert Decimal(str(api_trade["net_pnl"])) == expected_net
    assert Decimal(str(api_trade["closing_equity"])) == expected_closing_equity
    assert api_data["daily_compounding"][0]["net_pnl"] == float(expected_net)
    assert api_data["daily_compounding"][0]["closing_equity"] == float(expected_closing_equity)

    csv_response = export_ledger(format="csv", x_zerqen_dashboard_token="test-token")
    csv_rows = list(csv.DictReader(io.StringIO(csv_response.body.decode())))
    assert len(csv_rows) == 1
    assert Decimal(csv_rows[0]["Net P&L"]) == expected_net

    xlsx_response = export_ledger(format="xlsx", x_zerqen_dashboard_token="test-token")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_response.body), read_only=True, data_only=True)
    assert wb.sheetnames == [
        "Daily Compounding",
        "Trade Ledger",
        "Orders",
        "Fills",
        "Positions",
        "P&L",
        "Risk Events",
        "Equity Snapshots",
        "Audit Log",
    ]
    rows = list(wb["Trade Ledger"].iter_rows(values_only=True))
    assert len(rows) == 2
    headers = list(rows[0])
    net_idx = headers.index("net_pnl")
    assert Decimal(str(rows[1][net_idx])) == expected_net

    # Simulate a server restart by closing and reopening a new PostgreSQL connection.
    with psycopg.connect(clean_paper_db) as conn:
        persisted = conn.execute(
            "SELECT COUNT(*), COALESCE(SUM(net_pnl),0) FROM zerqen_paper_trades WHERE account_id='default'"
        ).fetchone()
        assert persisted[0] == 1
        assert Decimal(str(persisted[1])) == expected_net


def test_rejected_signal_is_persisted_with_exact_reason(clean_paper_db, monkeypatch):
    _deterministic_market(monkeypatch)
    status, payload = _request(
        monkeypatch,
        "initialize",
        starting_capital="10000",
        exchange_id="okx",
        symbol="BTC/USDT",
        timeframe="1h",
    )
    assert status == 200
    status, payload = _request(monkeypatch, "cycle")
    assert status == 200
    assert payload["live_trading"] is False
    with psycopg.connect(clean_paper_db) as conn:
        row = conn.execute(
            "SELECT signal_id,symbol,strategy,risk_decision,rejected,rejection_reason FROM zerqen_paper_decisions ORDER BY signal_timestamp DESC LIMIT 1"
        ).fetchone()
    assert row[1] == "BTC/USDT"
    assert row[2] == "baseline_trend"
    assert row[3] == "REJECTED"
    assert row[4] is True
    assert row[5]


def test_ledger_ui_is_database_api_driven():
    html = (ROOT / "index.html").read_text(encoding="utf-8")
    assert 'data-view="ledger"' in html
    assert 'data-view="compounding"' in html
    assert 'data-view="audit"' in html
    assert '"/api/ledger?' in html
    assert '"/api/ledger/export?format="' in html
    assert "PAPER_TEST_HARNESS" not in html
